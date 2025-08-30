from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

menu = Document("./menu.docx")

menu.tables[0].rows[0].cells[1].text = "Buffet varié"
menu.tables[0].rows[1].cells[1].text = "Palette � la Diable ou Colin � la Proven�al & Ebly"
menu.tables[0].rows[2].cells[1].text = "Fromage - Yaourt aromatisé"
print(menu.tables[0].rows[1].cells[0].text)

menu.save("menu2.docx")