from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.image import Inches

menu = Document("./template.docx")

days = {
    'lundi': 1,
    'mardi': 2,
    'mercredi': 3,
    'jeudi': 4,
    'vendredi': 5
}

def set_day(day, entrees, plats, desserts):
    result = [entrees, plats, desserts]
    for i in range(1, 3 + 1):
        for j in result[i - 1]:
            p = menu.tables[0].rows[i].cells[days[day]].add_paragraph(j['name'])
            if j['made'] == True:
                pic = p.add_run().add_picture('./made.png', width=Inches(0.2), height=Inches(0.2))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

set_day(
    'lundi', 
    [{ 'name': "Salade", 'made': False }],
     [{ 'name': 'Frites', 'made': False },
      { 'name': 'Pizza', 'made': False }], 
    [{ 'name': 'Glace', 'made': False }]
)

set_day(
    'mardi',
    [{ 'name': "Patates", 'made': False }], 
    [{ 'name': 'Burger', 'made': True }, 
     { 'name': 'Pizza', 'made': False }], 
    [{ 'name': "Sneakers", 'made': False }]
)

set_day(
    'mercredi', 
    [{ 'name': "Carottes", 'made': False }], 
    [{ 'name': "Frites", 'made': False }, { 'name': 'Burger', 'made': True }], 
    [{ 'name': "Donut", 'made': False }]
)

menu.paragraphs[0].text = "Menu de la cantine à la date du 26 mai 2025 au 30 mai 2025"

menu.save("menu.docx")